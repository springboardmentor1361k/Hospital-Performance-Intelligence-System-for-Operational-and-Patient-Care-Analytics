import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook

# Define constants
INPUT_DIR = "cleaned_data"
ADMISSION_FILE = f"{INPUT_DIR}/admission_cleaned.csv"
DEPARTMENT_FILE = f"{INPUT_DIR}/department_cleaned.csv"
OUTPUT_PATH = "hospital_final_dataset.xlsx"
READMISSION_WINDOW_DAYS = 30

# Load datasets
admission_df = pd.read_csv(ADMISSION_FILE)
department_df = pd.read_csv(DEPARTMENT_FILE)

# Convert dates to datetime objects
admission_df["admission_date"] = pd.to_datetime(admission_df["admission_date"])
admission_df["discharge_date"] = pd.to_datetime(admission_df["discharge_date"])

# Merge department names for readability
admission_df = admission_df.merge(
    department_df[["department_id", "department_name"]],
    on="department_id",
    how="left"
)

# Calculate Average Length of Stay (Overall & by Department)
overall_avg_los = admission_df["length_of_stay"].mean()

avg_los_by_department = (
    admission_df.groupby("department_name")["length_of_stay"]
    .agg(average_length_of_stay="mean", total_admissions="count")
    .reset_index()
    .sort_values("average_length_of_stay", ascending=False)
)
avg_los_by_department["average_length_of_stay"] = avg_los_by_department["average_length_of_stay"].round(2)

# Calculate Readmission Rate
# Sort chronologically per patient to find readmission gaps
admission_df = admission_df.sort_values(["patient_id", "admission_date"]).reset_index(drop=True)

# Calculate days since previous discharge
admission_df["previous_discharge_date"] = admission_df.groupby("patient_id")["discharge_date"].shift(1)
admission_df["days_since_prior_discharge"] = (
    admission_df["admission_date"] - admission_df["previous_discharge_date"]
).dt.days

# Flag valid readmissions within the 30-day window
admission_df["is_readmission"] = (
    (admission_df["days_since_prior_discharge"] >= 0) & 
    (admission_df["days_since_prior_discharge"] <= READMISSION_WINDOW_DAYS)
)

# Calculate overall rates
total_admissions = len(admission_df)
total_readmissions = admission_df["is_readmission"].sum()
overall_readmission_rate = (total_readmissions / total_admissions) * 100

# Calculate department-level readmission rates
readmission_by_department = (
    admission_df.groupby("department_name")
    .agg(
        total_admissions=("admission_id", "count"),
        total_readmissions=("is_readmission", "sum")
    )
    .reset_index()
)
readmission_by_department["readmission_rate_pct"] = (
    (readmission_by_department["total_readmissions"] / readmission_by_department["total_admissions"]) * 100
).round(2)
readmission_by_department = readmission_by_department.sort_values("readmission_rate_pct", ascending=False)

# Compile summary DataFrames
kpi_summary = pd.DataFrame({
    "KPI": ["Average Length of Stay (days)", f"Readmission Rate ({READMISSION_WINDOW_DAYS}-day, %)"],
    "Value": [round(overall_avg_los, 2), round(overall_readmission_rate, 2)]
})

admission_detail = admission_df[[
    "admission_id", "patient_id", "department_name", "admission_date", 
    "discharge_date", "length_of_stay", "days_since_prior_discharge", 
    "is_readmission"
]].copy()

# Export to Excel
with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
    kpi_summary.to_excel(writer, sheet_name="KPI_Summary", index=False)
    avg_los_by_department.to_excel(writer, sheet_name="Avg_LOS_by_Department", index=False)
    readmission_by_department.to_excel(writer, sheet_name="Readmission_by_Department", index=False)
    admission_detail.to_excel(writer, sheet_name="Admission_KPI_Detail", index=False)

# Apply formatting to Excel workbook
wb = load_workbook(OUTPUT_PATH)
header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    
    # Format headers
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    
    # Auto-adjust column widths
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 40)

wb.save(OUTPUT_PATH)
print(f"Data processing complete. Results exported to {OUTPUT_PATH}")